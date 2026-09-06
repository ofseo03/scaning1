"""Microsoft Word(.docx) 출력."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ..config import ConvertOptions
from ..errors import WriterError
from ..model import BlockKind, Document, Page
from .base import DocumentWriter

#: 제목 단계별 글자 크기 배율
HEADING_SCALE = {1: 1.65, 2: 1.35, 3: 1.15}


class DocxWriter(DocumentWriter):
    name = "docx"
    extension = ".docx"
    description = "Microsoft Word 문서"

    def write(self, document: Document, path: Path, options: ConvertOptions) -> Path:
        docx = DocxDocument()
        _setup_page(docx, options)
        _setup_default_style(docx, options)

        if document.title:
            heading = docx.add_paragraph()
            run = heading.add_run(document.title)
            _apply_font(run, options, size=options.font_size * 1.8, bold=True)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            heading.paragraph_format.space_after = Pt(18)

        for index, page in enumerate(document.pages):
            if index > 0 and options.page_break:
                docx.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            _write_page(docx, page, options)

        try:
            docx.save(str(path))
        except OSError as exc:
            raise WriterError(f"Word 문서를 저장하지 못했습니다: {path} ({exc})") from exc
        return path


def _write_page(docx, page: Page, options: ConvertOptions) -> None:
    if options.embed_image and page.image_path:
        try:
            docx.add_picture(page.image_path, width=Cm(15))
            docx.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:  # 이미지 삽입 실패가 변환 전체를 막지 않게 한다
            pass

    for block in page.blocks:
        if block.is_empty():
            continue
        if block.kind is BlockKind.HEADING:
            _add_heading(docx, block, options)
        elif block.kind is BlockKind.LIST_ITEM:
            _add_list_item(docx, block, options)
        else:
            _add_paragraph(docx, block, options)


def _add_heading(docx, block, options: ConvertOptions):
    paragraph = docx.add_paragraph()
    scale = HEADING_SCALE.get(block.level, 1.15)
    run = paragraph.add_run(block.text)
    _apply_font(run, options, size=options.font_size * scale, bold=True)
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(12)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.2
    return paragraph


def _add_list_item(docx, block, options: ConvertOptions):
    style_base = "List Number" if block.ordered else "List Bullet"
    style = style_base if block.level <= 0 else f"{style_base} {min(3, block.level + 1)}"
    try:
        paragraph = docx.add_paragraph(style=style)
    except KeyError:
        paragraph = docx.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.75 * (block.level + 1))
        block_text = f"{block.list_marker} {block.text}".strip()
        run = paragraph.add_run(block_text)
        _apply_font(run, options)
        return paragraph
    run = paragraph.add_run(block.text)
    _apply_font(run, options)
    paragraph.paragraph_format.line_spacing = options.line_spacing
    return paragraph


def _add_paragraph(docx, block, options: ConvertOptions):
    paragraph = docx.add_paragraph()
    for index, line in enumerate(block.text.split("\n")):
        run = paragraph.add_run(line)
        _apply_font(run, options)
        if index < block.text.count("\n"):
            run.add_break()
    fmt = paragraph.paragraph_format
    fmt.line_spacing = options.line_spacing
    fmt.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


# --- 서식 도우미 ---------------------------------------------------------

def _apply_font(run, options: ConvertOptions, *, size: float | None = None, bold: bool = False):
    """한글 글꼴이 제대로 적용되도록 eastAsia 속성까지 직접 지정한다."""
    run.font.size = Pt(size if size is not None else options.font_size)
    run.font.bold = bold
    run.font.name = options.font_latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), options.font_latin)
    rfonts.set(qn("w:hAnsi"), options.font_latin)
    rfonts.set(qn("w:eastAsia"), options.font_korean)
    rfonts.set(qn("w:cs"), options.font_latin)


def _setup_page(docx, options: ConvertOptions) -> None:
    section = docx.sections[0]
    section.page_width = Cm(21.0)     # A4
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.5)


def _setup_default_style(docx, options: ConvertOptions) -> None:
    style = docx.styles["Normal"]
    style.font.name = options.font_latin
    style.font.size = Pt(options.font_size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), options.font_latin)
    rfonts.set(qn("w:hAnsi"), options.font_latin)
    rfonts.set(qn("w:eastAsia"), options.font_korean)
    style.paragraph_format.line_spacing = options.line_spacing
